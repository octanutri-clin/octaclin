"""Verifica artefatos oficiais locais e prepara grade temporária para Jest.
Requer openpyxl apenas na pesquisa, não adiciona dependência ao produto.
Tolerância prospectiva de quantis publicados com 3 casas: 0,0005 unidade + 1e-10.
Não copia datasets completos ao Git ou runtime; todos os casos são referências públicas.
"""
import argparse, hashlib, json, pathlib, re
import openpyxl
ROOT=pathlib.Path(__file__).resolve().parents[3]
a=argparse.ArgumentParser();a.add_argument('--fontes',type=pathlib.Path,required=True);a.add_argument('--output',type=pathlib.Path,required=True);args=a.parse_args()
linhas=[l for l in (ROOT/'docs/product/FONTES_ANTROPOMETRIA_FASE_314.md').read_text().splitlines() if '[planilha]' in l]
rows=[];artifacts=[]
for n,line in enumerate(linhas):
    p=args.fontes/f'oms-planilha-{n:02}.xlsx'
    if not p.exists():continue
    esperado=re.findall(r'`([a-f0-9]{64})`',line)[0]
    atual=hashlib.sha256(p.read_bytes()).hexdigest()
    assert atual==esperado,('hash divergente',n)
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    vals=list(w.active.values);head=vals.pop(0)
    assert len(w.sheetnames)==1
    expected_sheet=re.findall('`([^`]+)`',line)[0]
    assert w.sheetnames[0]==expected_sheet
    columns=[]
    for column,name in enumerate(head):
        match=re.fullmatch(r'SD(\d+)(neg)?',str(name))
        if match: columns.append((column,-int(match.group(1)) if match.group(2) else int(match.group(1))))
    assert len(columns)>=9,head
    for r in vals:
        assert all(isinstance(v,(float,int)) for v in r)
        assert r[2]>0 and r[3]>0
        rows.append({'arquivo':n,'sheet':w.sheetnames[0],'x':r[0], 'l':r[1],'m':r[2],'s':r[3],
                     'caudas':not ('LFA_' in w.sheetnames[0] or 'hfa_' in w.sheetnames[0]),
                     'escores':[z for _,z in columns],'quantis':[r[column] for column,_ in columns]})
    artifacts.append({'arquivo':n,'sha256':atual,'sheet':w.sheetnames[0],'linhas':len(vals),'minimo':vals[0][0],'maximo':vals[-1][0]})
summary={'tolerancia_medida':.0005+1e-10,'openpyxl':openpyxl.__version__,'arquivos':artifacts,'linhas_total':len(rows),'grade':rows}
args.output.write_text(json.dumps(summary)+'\n')
print(json.dumps({k:v for k,v in summary.items() if k!='grade'},indent=2))
